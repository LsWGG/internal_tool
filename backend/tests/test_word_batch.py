import io
import json
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.word_batch_engine import checked_zip, read_excel, render_document, template_fields
from app.word_batch_manager import WordBatchManager
from app.word_batch_api import make_router


def image_bytes(color='blue'):
    output = io.BytesIO()
    Image.new('RGB', (120, 60), color).save(output, 'PNG')
    return output.getvalue()


def template_bytes():
    doc = Document()
    doc.add_heading('入职通知', 0)
    p = doc.add_paragraph()
    p.add_run('您好，').bold = True
    p.add_run('{{姓').italic = True
    p.add_run('名}}，编号 {{编号}}。')
    doc.add_paragraph('{{姓名}} / {{姓名}} / {{备注}}')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '姓名'
    table.cell(0, 1).text = '{{姓名}}'
    table.cell(1, 0).text = '照片'
    table.cell(1, 1).text = '{{%照片}}'
    doc.sections[0].header.paragraphs[0].text = '部门 {{部门}}'
    doc.sections[0].footer.paragraphs[0].text = '{{姓名}} 的文档'
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def excel_bytes(embedded=False, missing=False):
    wb = Workbook()
    ws = wb.active
    ws.title = '员工名单'
    ws.append(['姓名', '编号', '备注', '照片', '部门'])
    ws.append(['张三', 12, '第一行\n第二行 <&>', 'photo.png', '研发部'])
    ws.append(['张三', 0, '', 'missing.png' if missing else 'photo.png', '设计部'])
    ws['B2'].number_format = '00000'
    if embedded:
        ws['D2'] = None
        ws.add_image(ExcelImage(io.BytesIO(image_bytes('red'))), 'D2')
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


class EngineTests(unittest.TestCase):
    def test_split_runs_tables_headers_and_images(self):
        source = template_bytes()
        fields = template_fields(source)
        self.assertEqual({f['key'] for f in fields}, {'姓名', '编号', '备注', '照片', '部门'})
        mapping = {f['key']: {'column': f['key'], 'type': 'image' if f['image'] else 'text', 'width_mm': 40, 'height_mm': 40} for f in fields}
        values = {'姓名': '张三', '编号': '00012', '备注': '第一行\n第二行 <&>', '照片': 'photo.png', '部门': '研发部'}
        result = render_document(source, values, mapping, lambda _: image_bytes())
        doc = Document(io.BytesIO(result))
        self.assertEqual(doc.paragraphs[1].text, '您好，张三，编号 00012。')
        self.assertTrue(doc.paragraphs[1].runs[1].italic)
        self.assertEqual(doc.paragraphs[2].text, '张三 / 张三 / 第一行\n第二行 <&>')
        self.assertEqual(doc.tables[0].cell(0, 1).text, '张三')
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, '部门 研发部')
        self.assertEqual(doc.sections[0].footer.paragraphs[0].text, '张三 的文档')
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertAlmostEqual(doc.inline_shapes[0].width / doc.inline_shapes[0].height, 2)
        self.assertEqual(template_bytes(), source)
        with zipfile.ZipFile(io.BytesIO(result)) as archive:
            self.assertFalse(any(b'{{' in archive.read(n) for n in archive.namelist() if n.endswith('.xml')))

    def test_embedded_images_and_number_formats(self):
        images = {}
        sheets = read_excel(excel_bytes(embedded=True), lambda name, data: images.update({name: data}))
        self.assertEqual(sheets[0]['rows'][0]['values']['编号'], '00012')
        self.assertEqual(sheets[0]['rows'][1]['values']['编号'], '0')
        self.assertIn(sheets[0]['rows'][0]['values']['照片'], images)
        self.assertEqual(sheets[0]['image_columns'], ['照片'])

    def test_duplicate_headers_and_unsafe_zip(self):
        wb = Workbook()
        wb.active.append(['姓名', '姓名'])
        wb.active.append(['a', 'b'])
        output = io.BytesIO()
        wb.save(output)
        with self.assertRaisesRegex(ValueError, '重复表头'):
            read_excel(output.getvalue(), lambda *_: None)
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as archive:
            archive.writestr('../escape.png', b'bad')
        with self.assertRaisesRegex(ValueError, '不安全路径'):
            checked_zip(output.getvalue())


class ManagerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.manager = WordBatchManager(self.temp.name)

    def tearDown(self):
        self.manager.pool.shutdown(wait=True)
        self.temp.cleanup()

    def create_workspace(self, missing=False):
        space = self.manager.prepare('通知.docx', template_bytes(), '数据.xlsx', excel_bytes(missing=missing), [('photo.png', image_bytes())])
        config = {'sheet': '员工名单', 'filename': '{{姓名}}', 'mappings': {f['key']: {'column': f['key'], 'type': 'image' if f['image'] else 'text'} for f in space['fields']}}
        return space, config

    def wait_task(self, ident):
        deadline = time.monotonic() + 15
        while time.monotonic() < deadline:
            with self.manager.lock:
                task = dict(self.manager.tasks[ident])
            if task['status'] not in ('queued', 'running'):
                return task
            time.sleep(.02)
        self.fail('Task timed out')

    def test_full_flow_and_cleanup(self):
        space, config = self.create_workspace()
        name = self.manager.preview(space['id'], config)
        self.assertTrue((self.manager.path(space['id']) / name).is_file())
        task = self.manager.create(space['id'], config)
        self.manager.delete(space['id'], True)
        completed = self.wait_task(task['id'])
        self.assertEqual(completed['status'], 'completed', completed)
        names = [r['name'] for r in completed['results']]
        self.assertEqual(names, ['张三.docx', '张三_2.docx'])
        with zipfile.ZipFile(self.manager.result_file(task['id'])) as archive:
            self.assertEqual(archive.namelist(), names)
        self.manager.delete(task['id'])
        self.assertEqual(list(Path(self.temp.name).iterdir()), [])

    def test_partial_failure_retry_and_restart(self):
        space, config = self.create_workspace(missing=True)
        task = self.manager.create(space['id'], config)
        completed = self.wait_task(task['id'])
        self.assertEqual(completed['status'], 'partial', completed)
        self.assertEqual(completed['success'], 1)
        self.assertIn('未找到图片', completed['results'][1]['error'])
        self.manager.retry(task['id'])
        self.assertEqual(self.wait_task(task['id'])['success'], 1)
        restored = WordBatchManager(self.temp.name)
        try:
            self.assertEqual(restored.list_tasks()[0]['id'], task['id'])
        finally:
            restored.pool.shutdown(wait=True)

    def test_validation(self):
        space, config = self.create_workspace()
        config['mappings']['照片']['width_mm'] = 999
        with self.assertRaisesRegex(ValueError, '最大宽度'):
            self.manager.create(space['id'], config)
        with self.assertRaises(ValueError):
            self.manager.delete('../')


class ApiTests(unittest.TestCase):
    def test_matching_downloadable_examples(self):
        with tempfile.TemporaryDirectory() as root:
            app = FastAPI()
            app.include_router(make_router(root))
            with TestClient(app) as client:
                word = client.get('/api/word-batch/examples/word')
                excel = client.get('/api/word-batch/examples/excel')
                self.assertEqual(word.status_code, 200)
                self.assertEqual(excel.status_code, 200)
                self.assertIn('attachment', word.headers['content-disposition'])
                self.assertEqual(client.get('/api/word-batch/examples/unknown').status_code, 404)
                self.assertFalse(list(Path(root).iterdir()), 'Examples should not create task folders')
                images = {}
                sheets = read_excel(excel.content, lambda name, data: images.update({name: data}))
                fields = template_fields(word.content)
                self.assertEqual(len(sheets[0]['rows']), 3)
                self.assertEqual(len(images), 3)
                self.assertEqual({f['key'] for f in fields}, set(sheets[0]['columns']))
                mapping = {f['key']: {'column': f['key'], 'type': 'image' if f['image'] else 'text', 'width_mm': 40, 'height_mm': 40} for f in fields}
                for row in sheets[0]['rows']:
                    result = render_document(word.content, row['values'], mapping, images.__getitem__)
                    document = Document(io.BytesIO(result))
                    self.assertEqual(len(document.inline_shapes), 1)
                    self.assertEqual(document.tables[0].cell(1, 1).text, row['values']['员工编号'])
                    self.assertNotIn('{{', document.element.xml)

    def test_upload_preview_create_export_and_delete(self):
        with tempfile.TemporaryDirectory() as root:
            app = FastAPI()
            app.include_router(make_router(root))
            with TestClient(app) as client:
                response = client.post('/api/word-batch/workspaces', files=[('template', ('test.docx', template_bytes())), ('excel', ('data.xlsx', excel_bytes())), ('images', ('photo.png', image_bytes()))])
                self.assertEqual(response.status_code, 200, response.text)
                space = response.json()
                config = {'sheet': '员工名单', 'mappings': {f['key']: {'column': f['key'], 'type': 'image' if f['image'] else 'text'} for f in space['fields']}}
                preview = client.post(f'/api/word-batch/workspaces/{space["id"]}/preview', json=config)
                self.assertEqual(preview.status_code, 200, preview.text)
                self.assertTrue(client.get(preview.json()['url']).content.startswith(b'PK'))
                response = client.post(f'/api/word-batch/workspaces/{space["id"]}/tasks', json=config)
                self.assertEqual(response.status_code, 200, response.text)
                ident = response.json()['id']
                deadline = time.monotonic() + 15
                while time.monotonic() < deadline:
                    task = client.get('/api/word-batch/tasks').json()[0]
                    if task['status'] not in ('running', 'queued'):
                        break
                    time.sleep(.02)
                self.assertEqual(task['status'], 'completed', task)
                self.assertEqual(client.get(f'/api/word-batch/tasks/{ident}/results').json()['total'], 2)
                self.assertTrue(client.get(f'/api/word-batch/tasks/{ident}/export').content.startswith(b'PK'))
                self.assertTrue(client.get(f'/api/word-batch/tasks/{ident}/files/1').content.startswith(b'PK'))
                self.assertEqual(client.delete(f'/api/word-batch/tasks/{ident}').status_code, 200)
                self.assertEqual(client.delete(f'/api/word-batch/workspaces/{space["id"]}').status_code, 200)
                self.assertFalse(list(Path(root).iterdir()))


if __name__ == '__main__':
    unittest.main()
