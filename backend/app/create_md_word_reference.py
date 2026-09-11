from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def font(style, name, size, color=None, bold=None, east_asia="Arial Unicode MS"):
    style.font.name = name
    style.font.size = Pt(size)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def build(path: Path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.82)
    section.left_margin = section.right_margin = Inches(.9)
    section.header_distance = section.footer_distance = Inches(.42)
    normal = doc.styles["Normal"]
    font(normal, "Aptos", 10.5, "26364D")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color, before, after in (
        ("Title", 27, "173B6C", 0, 12), ("Heading 1", 18, "245FA8", 18, 8),
        ("Heading 2", 14, "315FC8", 14, 6), ("Heading 3", 11.5, "4C6688", 10, 4)):
        style = doc.styles[name]
        font(style, "Aptos Display" if name != "Heading 3" else "Aptos", size, color, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Code" not in [style.name for style in doc.styles]:
        doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    font(doc.styles["Code"], "Cascadia Mono", 9, "24364C")
    doc.styles["Code"].paragraph_format.space_after = Pt(5)
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        font(style, "Aptos", 10.5, "26364D")
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.18
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for cell in table.rows[0].cells:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E9F1FB")
        cell._tc.get_or_add_tcPr().append(shading)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Markdown → Word")
    run.font.name, run.font.size, run.font.color.rgb = "Aptos", Pt(8), RGBColor(112, 131, 154)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    doc.save(path)


if __name__ == "__main__":
    target = Path(__file__).resolve().parent / "assets" / "md_word_reference.docx"
    target.parent.mkdir(parents=True, exist_ok=True)
    build(target)
    print(target)
