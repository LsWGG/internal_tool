"""Literal DOCX mail merge: no Python/Jinja expressions or remote image fetching."""
import io
import re
import zipfile
from datetime import date, datetime, time
from pathlib import PurePosixPath
from types import SimpleNamespace

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from PIL import Image, ImageOps

TOKEN = re.compile(r"\{\{\s*([^{}\r\n]+?)\s*\}\}")
LIMIT = 200 * 1024 * 1024
IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff', '.heic', '.heif'}


def safe_name(name):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(name)).strip(' .')[:120]
    return name if name and name not in {'.', '..'} else '文档'


def checked_zip(content):
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
        entries = archive.infolist()
        if len(entries) > 5000 or sum(i.file_size for i in entries) > LIMIT:
            raise ValueError('解压后文件过大或文件数量过多（最多 200MB / 5000 项）')
        for item in entries:
            path = PurePosixPath(item.filename.replace('\\', '/'))
            if path.is_absolute() or '..' in path.parts or (item.external_attr >> 16) & 0o170000 == 0o120000:
                raise ValueError('压缩包含不安全路径或符号链接')
        return archive
    except zipfile.BadZipFile as exc:
        raise ValueError('文件不是有效的 DOCX / XLSX / ZIP 文件') from exc


def open_template(content):
    with checked_zip(content) as archive:
        if 'word/document.xml' not in archive.namelist():
            raise ValueError('请选择有效的 .docx Word 模板')
        for entry in archive.infolist():
            if 'vbaProject' in entry.filename or '/embeddings/' in entry.filename:
                raise ValueError('模板不能包含宏或嵌入式程序对象')
            if entry.filename.endswith('.rels'):
                from lxml import etree
                root = etree.fromstring(archive.read(entry), etree.XMLParser(resolve_entities=False, no_network=True))
                for rel in root:
                    if rel.get('TargetMode') == 'External' and not rel.get('Type', '').endswith('/hyperlink'):
                        raise ValueError('模板包含外链图片或外部数据，请先在 Word 中嵌入内容')
    return Document(io.BytesIO(content))


def paragraphs(document):
    parts = [document.part] + [part for part in document.part.package.parts
                              if str(part.partname).startswith(('/word/header', '/word/footer'))]
    for part in parts:
        for element in part.element.xpath('.//w:p'):
            # Text boxes have nested paragraphs; do not process their text twice.
            nodes = [n for n in element.xpath('.//w:t')
                     if next(n.iterancestors(qn('w:p')), None) is element]
            yield Paragraph(element, SimpleNamespace(part=part)), nodes


def template_fields(content):
    fields = {}
    for _, nodes in paragraphs(open_template(content)):
        text = ''.join(n.text or '' for n in nodes)
        for token in TOKEN.finditer(text):
            key = token.group(1).strip()
            image = key.startswith('%')
            key = key.lstrip('%').strip()
            if not key or len(key) > 100:
                raise ValueError('模板字段名称为空或超过 100 个字符')
            fields[key] = {'key': key, 'image': image or fields.get(key, {}).get('image', False)}
    if not fields:
        raise ValueError('未发现占位符。请在模板正文、表格或页眉页脚中填写 {{姓名}} 等字段')
    if len(fields) > 200:
        raise ValueError('模板字段不能超过 200 个')
    return list(fields.values())


def normalized_image(content):
    try:
        try:
            from pillow_heif import register_heif_opener
            register_heif_opener()
        except ImportError:
            pass
        with Image.open(io.BytesIO(content)) as original:
            if original.width * original.height > 40_000_000:
                raise ValueError('图片超过 4000 万像素，请先缩小')
            image = ImageOps.exif_transpose(original)
            image.thumbnail((2400, 2400), Image.Resampling.LANCZOS)
            output = io.BytesIO()
            image.convert('RGBA' if 'A' in image.getbands() else 'RGB').save(output, format='PNG')
            return output.getvalue()
    except Exception as exc:
        raise ValueError(f'无法读取图片：{exc}') from exc


def display_cell(cell):
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d' if value.time() == time() else '%Y-%m-%d %H:%M:%S')
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return '是' if value else '否'
    if isinstance(value, (float, int)):
        fmt = cell.number_format or ''
        if re.fullmatch('0+', fmt):
            return str(int(value)).zfill(len(fmt)) if float(value).is_integer() else str(value)
        if re.fullmatch(r'0(?:\.0+)?%', fmt):
            decimals = len(fmt.split('.')[1].rstrip('%')) if '.' in fmt else 0
            return f'{value * 100:.{decimals}f}%'
        if re.fullmatch(r'(?:#,##)?0\.0+', fmt):
            return f'{value:.{len(fmt.split(".")[1])}f}'
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value)


def read_excel(content, add_image):
    with checked_zip(content) as archive:
        if 'xl/workbook.xml' not in archive.namelist():
            raise ValueError('请选择 .xlsx Excel 文件，不支持旧版 .xls')
        if any('vbaProject' in name for name in archive.namelist()):
            raise ValueError('不接受包含宏的 Excel')
    workbook = load_workbook(io.BytesIO(content), data_only=True, keep_links=False)
    formulas = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    sheets = []
    try:
        for ws in workbook.worksheets:
            if ws.max_row > 5001 or ws.max_column > 200:
                raise ValueError(f'工作表“{ws.title}”超过 5000 行或 200 列，请拆分数据')
            headers = [str(cell.value).strip() if cell.value is not None else '' for cell in ws[1]]
            while headers and not headers[-1]:
                headers.pop()
            if not headers:
                continue
            if any(not name for name in headers) or len(set(headers)) != len(headers):
                raise ValueError(f'工作表“{ws.title}”第一行有空白或重复表头，请修正后上传')
            images, image_columns = {}, set()
            for index, image in enumerate(ws._images):
                anchor = getattr(image.anchor, '_from', None)
                if anchor is None or anchor.row < 1 or anchor.col >= len(headers):
                    raise ValueError(f'工作表“{ws.title}”的图片须锚定到数据行的有效列')
                position = (anchor.row + 1, anchor.col)
                if position in images:
                    raise ValueError(f'工作表“{ws.title}”第 {position[0]} 行同一单元格存在多张图片')
                name = f'__excel_{len(sheets)}_{index}.png'
                add_image(name, image._data())
                images[position] = name
                image_columns.add(headers[anchor.col])
            rows, warnings = [], []
            last_row = max(ws.max_row, max((p[0] for p in images), default=0))
            if last_row > 5001:
                raise ValueError('图片锚定的行超出 5000 行数据范围')
            for row in ws.iter_rows(min_row=2, max_row=last_row, max_col=len(headers)):
                values = {}
                for col, cell in enumerate(row):
                    formula = formulas[ws.title].cell(cell.row, col + 1)
                    if formula.data_type == 'f' and cell.value is None:
                        if len(warnings) < 10:
                            warnings.append(f'{cell.coordinate} 公式无缓存结果，请在 Excel 中重新计算并保存')
                    values[headers[col]] = images.get((cell.row, col), display_cell(cell))
                if any(value != '' for value in values.values()):
                    rows.append({'row': row[0].row, 'values': values})
            if rows:
                sheets.append({'name': ws.title, 'columns': headers, 'rows': rows,
                               'image_columns': sorted(image_columns), 'warnings': warnings})
    finally:
        workbook.close()
        formulas.close()
    if not sheets:
        raise ValueError('Excel 没有数据行：第一行为表头，第二行开始填写数据')
    return sheets


def render_document(template, values, mappings, image_loader):
    document = open_template(template)
    for paragraph, nodes in paragraphs(document):
        text = ''.join(node.text or '' for node in nodes)
        spans, offset = [], 0
        for node in nodes:
            size = len(node.text or '')
            spans.append((offset, offset + size, node))
            offset += size
        for match in reversed(list(TOKEN.finditer(text))):
            key = match.group(1).strip().lstrip('%').strip()
            mapping = mappings[key]
            value = values.get(mapping['column'], '')
            start, end = match.span()
            selected = [(a, b, node) for a, b, node in spans if b > start and a < end]
            first_a, _, first = selected[0]
            last_a, _, last = selected[-1]
            prefix = (first.text or '')[:start - first_a]
            suffix = (last.text or '')[end - last_a:]
            for _, _, node in selected:
                node.text = ''
            first.text = prefix
            first.set(qn('xml:space'), 'preserve')
            inserts = []
            if mapping['type'] == 'image' and value:
                data = image_loader(str(value))
                with Image.open(io.BytesIO(data)) as image:
                    width, height = image.size
                box_w, box_h = mapping['width_mm'], mapping['height_mm']
                ratio = min(box_w / width, box_h / height)
                run = paragraph.add_run()
                picture = run.add_picture(io.BytesIO(data), width=Mm(width * ratio), height=Mm(height * ratio))
                # Move only the drawing into the original styled run at the token position.
                inserts.append(picture._inline.getparent())
                paragraph._p.remove(run._r)
            elif value:
                for index, line in enumerate(str(value).replace('\r\n', '\n').replace('\r', '\n').split('\n')):
                    if index:
                        inserts.append(OxmlElement('w:br'))
                    node = OxmlElement('w:t')
                    node.set(qn('xml:space'), 'preserve')
                    node.text = line
                    inserts.append(node)
            if first is last:
                node = OxmlElement('w:t')
                node.set(qn('xml:space'), 'preserve')
                node.text = suffix
                inserts.append(node)
            else:
                last.text = suffix
                last.set(qn('xml:space'), 'preserve')
            previous = first
            for node in inserts:
                previous.addnext(node)
                previous = node
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
