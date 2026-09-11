"""Small matching examples generated in memory; no user task files are created."""
import io
from datetime import date
from functools import lru_cache

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw


@lru_cache(maxsize=2)
def example_file(kind):
    output = io.BytesIO()
    if kind == 'word':
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(22)
        section.left_margin = section.right_margin = Mm(25)
        for name in ('Normal', 'Title'):
            style = document.styles[name]
            style.font.name = 'Arial'
            style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
            style.font.color.rgb = RGBColor(0, 0, 0)
        document.styles['Normal'].font.size = Pt(11)
        document.styles['Normal'].paragraph_format.space_after = Pt(8)
        document.styles['Title'].font.size = Pt(22)
        document.add_paragraph('员工信息卡', style='Title')
        document.add_paragraph('本信息卡根据员工名单生成，汇总员工的部门、岗位、入职日期和照片。')
        table = document.add_table(rows=0, cols=2)
        table.autofit = False
        table.columns[0].width, table.columns[1].width = Mm(35), Mm(125)
        for label, value in [('姓名', '{{姓名}}'), ('员工编号', '{{员工编号}}'), ('部门', '{{部门}}'),
                             ('岗位', '{{岗位}}'), ('入职日期', '{{入职日期}}'), ('照片', '{{%照片}}'), ('备注', '{{备注}}')]:
            cells = table.add_row().cells
            for index, text in enumerate((label, value)):
                cells[index].text = text
                tcpr = cells[index]._tc.get_or_add_tcPr()
                margins = OxmlElement('w:tcMar')
                for edge in ('top', 'left', 'bottom', 'right'):
                    item = OxmlElement('w:' + edge)
                    item.set(qn('w:w'), '120')
                    item.set(qn('w:type'), 'dxa')
                    margins.append(item)
                tcpr.append(margins)
                if index == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'F2F4F7')
                    tcpr.append(shade)
            if label == '照片':
                cells[1].paragraphs[0].alignment = 1
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            item = OxmlElement('w:' + edge)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                item.set(qn('w:' + key), value)
            borders.append(item)
        table._tbl.tblPr.append(borders)
        section.footer.paragraphs[0].text = '员工信息卡 · {{姓名}} · {{员工编号}}'
        document.save(output)
    elif kind == 'excel':
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '员工名单'
        sheet.append(['姓名', '员工编号', '部门', '岗位', '入职日期', '照片', '备注'])
        records = [
            ['张三', '00001', '研发部', '软件工程师', date(2026, 9, 1), None, '负责产品研发与维护。'],
            ['李四', '00002', '设计部', '产品设计师', date(2026, 9, 2), None, '负责界面设计。\n参与产品体验优化。'],
            ['王五', '00003', '运营部', '运营专员', date(2026, 9, 3), None, '负责日常运营与用户支持。'],
        ]
        for row, record in enumerate(records, 2):
            sheet.append(record)
            sheet.row_dimensions[row].height = 82
            sheet.cell(row, 2).number_format = '@'
            sheet.cell(row, 5).number_format = 'yyyy-mm-dd'
            picture = Image.new('RGB', (144, 144), ['#e8effb', '#e7f2ee', '#f2ecfa'][row - 2])
            draw = ImageDraw.Draw(picture)
            draw.ellipse((48, 24, 96, 72), fill='#7d91b2')
            draw.rounded_rectangle((28, 80, 116, 138), radius=26, fill='#7d91b2')
            source = io.BytesIO()
            picture.save(source, 'PNG')
            source.seek(0)
            image = ExcelImage(source)
            image.width = image.height = 96
            sheet.add_image(image, f'F{row}')
        widths = {'A': 14, 'B': 16, 'C': 16, 'D': 22, 'E': 18, 'F': 17, 'G': 40}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        for row in sheet:
            for cell in row:
                cell.font = Font(name='宋体', size=11, bold=cell.row == 1, color='FFFFFF' if cell.row == 1 else '24364B')
                cell.fill = PatternFill('solid', fgColor='344D70' if cell.row == 1 else 'F3F6FA' if cell.row % 2 == 0 else 'FFFFFF')
                cell.alignment = Alignment(horizontal='left' if cell.column in (4, 7) else 'center', vertical='center', wrap_text=True)
                cell.border = Border(*( [Side(style='thin', color='D9D9D9')] * 4))
        sheet.row_dimensions[1].height = 28
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = 'A1:G4'
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        workbook.save(output)
        workbook.close()
    else:
        raise ValueError('示例不存在')
    return output.getvalue()
